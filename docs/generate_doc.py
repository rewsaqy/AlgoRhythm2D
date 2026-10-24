import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = Document()

# Page Margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

# Color Palette
C_DARK = RGBColor(0x0F, 0x17, 0x2A)       # #0F172A Slate 900
C_PRIMARY = RGBColor(0x02, 0x84, 0xC7)    # #0284C7 Cyber Cyan/Blue
C_SECONDARY = RGBColor(0x0D, 0x94, 0x88)  # #0D9488 Teal
C_MUTED = RGBColor(0x64, 0x74, 0x8B)      # #64748B Slate 500
C_ALERT = RGBColor(0xDC, 0x26, 0x26)      # #DC2626 Red

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header(title, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.font.color.rgb = C_PRIMARY
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(11.5)
        run.font.color.rgb = C_DARK
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = C_SECONDARY
    return p

def add_body(text, bold_prefix=None, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Arial"
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = C_DARK
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = C_DARK
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.name = "Arial"
        r_b.font.size = Pt(9.5)
        r_b.font.color.rgb = C_DARK
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = C_DARK
    return p

# TITLE BLOCK
p_top = doc.add_paragraph()
p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_top.paragraph_format.space_after = Pt(2)
r_sub = p_top.add_run("OLIMPIADE TIK INFORMATIKA NASIONAL (OTN) VIII 2026\nKATEGORI LOMBA MEMBUAT GAME — TINGKAT SMA/SMK")
r_sub.font.name = "Arial"
r_sub.font.size = Pt(9)
r_sub.font.color.rgb = C_MUTED
r_sub.bold = True

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(4)
r_title = p_title.add_run("DOKUMEN TEKNIS & PENJELASAN KARYA\nAlgoRhythm: CODE IS SHIELD")
r_title.font.name = "Arial"
r_title.font.size = Pt(18)
r_title.bold = True
r_title.font.color.rgb = C_PRIMARY

p_tagline = doc.add_paragraph()
p_tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_tagline.paragraph_format.space_after = Pt(16)
r_tag = p_tagline.add_run("Game Rhythm-Action Hibrida dengan Otomasi Logika Rule-Engine dan Simulasi Keamanan Siber")
r_tag.font.name = "Arial"
r_tag.font.size = Pt(10)
r_tag.italic = True
r_tag.font.color.rgb = C_MUTED

# ==================== BAB 1 ====================
add_header("BAB I. KONSEP DAN LATAR BELAKANG KARYA", level=1)

add_header("1.1 Latar Belakang Masalah & Urgensi", level=2)
add_body("Di era transformasi digital dan pesatnya adopsi Artificial Intelligence, kemampuan computational thinking (berpikir komputasional) dan kesadaran keamanan siber (cybersecurity awareness) menjadi fondasi literasi paling krusial bagi generasi muda. Namun, dalam konteks pendidikan di sekolah, pembelajaran materi logika pemrograman, percabangan kondisi (if-else), dan arsitektur firewall sering kali disajikan secara teoretis, statis, dan berjarak dari pengalaman interaktif siswa.")
add_body("AlgoRhythm: CODE IS SHIELD lahir sebagai media edukatif inovatif yang mengawinkan genre rhythm game (permainan ritme ketukan) dengan rekayasa otomasi sistem pertahanan. Pendekatan game ini menempatkan kode pemrograman bukan sebagai teks pasif, melainkan sebagai perisai aktif yang melindungi integritas sistem dari gempuran paket data berbahaya.")

add_header("1.2 Narasi, Worldbuilding, dan Karakter A.I.D.A.", level=2)
add_body("Pemain berperan sebagai Cyber Defense Operator yang ditugaskan menjaga Core Node dari infeksi Adware dan serangan paket siber tak dikenal. Dalam menjalankan tugasnya, operator didampingi oleh A.I.D.A. (Automated Intelligence Defense Assistant), sebuah entitas AI sistem yang bertugas memantau telemetri, memberikan peringatan ancaman, dan menyajikan analisis pasca-serangan.")
add_body("Keberadaan A.I.D.A. mengejawantahkan sub-tema 'Bijak Menggunakan AI': AI dihadirkan bukan sebagai entitas otonom yang menggantikan kendali manusia, melainkan sebagai sistem pembantu cerdas (co-pilot). Efektivitas AI sangat bergantung pada bagaimana operator menyusun aturan logika, menentukan ambang batas toleransi, dan mengalokasikan sumber daya energi secara rasional.")

# ==================== BAB 2 ====================
add_header("BAB II. ARSITEKTUR GAMEPLAY DAN SISTEM MEKANIK", level=1)

add_header("2.1 Dual-Defense System (Sinergi Manual vs Otomasi)", level=2)
add_body("Mekanik inti AlgoRhythm mengadopsi sistem pertahanan ganda yang memaksa pemain melakukan continuous trade-off management di bawah tekanan tempo musik:")

add_bullet(" Pemain membaca telegraf visual ancaman yang meluncur dari 4 arah kardinal (Up, Down, Left, Right). Pemain harus menekan tombol perisai tepat pada audio beat target. Keberhasilan manual dengan akurasi tinggi menghasilkan rating PERFECT, mengembalikan HP (Integrity), serta mengisi ulang cadangan ENERGY sistem.", bold_prefix="1. Manual Rhythm Defense (Refleks & Ritme DSP):")
add_bullet(" Pemain dapat menyusun dan mengaktifkan skrip aturan logika pada Rule Engine (contoh: IF Threat.Direction == UP THEN AutoShield.Deploy(UP)). Saat paket malware terdeteksi sesuai kondisi, sistem otomatis menangkis serangan tanpa intervensi tombol arah manual. Otomasi ini mengonsumsi Energy dan penilaian dibatasi pada tingkat GOOD.", bold_prefix="2. Automated Rule-Engine Shield (Computational Automation):")
add_bullet(" Sistem otomasi sangat efektif menangani pola serangan burst cepat atau serangan simultan yang melebihi batas kecepatan fisik pemain. Namun, karena otomasi terus menguras energi dan tidak menghasilkan Perfect, pemain wajib melakukan pertahanan manual secara disiplin untuk regenerasi energi. Ini merefleksikan prinsip nyata optimasi algoritma: trade-off antara kecepatan, presisi, dan konsumsi resource.", bold_prefix="3. Strategic Trade-Off (Energy Loop):")

add_header("2.2 Taksonomi Ancaman Malware (Threat Variants)", level=2)
add_body("Pola serangan dirancang tidak seragam guna menguji logika adaptasi pemain:")
add_bullet(" Paket data berbahaya standar dengan kecepatan konstan yang mengikuti beat 1/1 lagu.", bold_prefix="• Adware Standard: ")
add_bullet(" Serangan paket ganda atau beruntun dengan interval cepat (off-beat 1/2 atau 1/4), sangat ideal ditangkis menggunakan Rule-Engine otomasi.", bold_prefix="• High-Speed Burst: ")
add_bullet(" Paket data palsu yang mengecoh. Jika pemain bertahan manual pada decoy, energi akan terbuang sia-sia; rule engine yang terprogram cerdas dapat dilatih untuk mengabaikan sinyal palsu ini.", bold_prefix="• Decoy / Glitch Threat: ")

add_header("2.3 Siklus Gelombang dan Fase Shop Terminal (Refactoring Phase)", level=2)
add_body("Struktur permainan dibagi menjadi beberapa Wave yang intens. Di antara transisi wave, permainan beralih ke Terminal Interface (Shop & Config):")
add_bullet(" Menggunakan kompensasi komputasi yang diperoleh untuk memulihkan kerusakan integrity sistem.", bold_prefix="• System Patch / Integrity Repair: ")
add_bullet(" Menambah kapasitas baterai firewall dan efisiensi konsumsi daya per eksekusi rule.", bold_prefix="• Memory & Energy Optimization: ")
add_bullet(" Membuka token kondisi baru (misalnya: filter kecepatan ancaman, trigger ambang darurat) untuk memperluas kapabilitas otomasi di wave berikutnya.", bold_prefix="• Rule Expansion Chips: ")

# ==================== BAB 3 ====================
add_header("BAB III. FITUR TEKNIS DAN KEUNGGULAN INOVASI", level=1)

# Table of Features
table = doc.add_table(rows=6, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

col_widths = [Inches(2.2), Inches(4.6)]
headers = ["Fitur Inovatif", "Deskripsi & Keunggulan Teknis"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.width = col_widths[i]
    set_cell_background(cell, "0284C7")
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(h)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

data_table = [
    ("DSP-Accurate Audio Timing", "Menggunakan AudioSettings.dspTime sebagai otoritas absolut waktu ritme. Menghilangkan ketergantungan pada frame-rate (Update delta-time) sehingga beat synchro tetap presisi mutlak tanpa audio drift di berbagai spesifikasi hardware PC."),
    ("Modular Safe Rule-Engine", "Sistem eksekusi kondisi-aksi yang dirancang dengan prinsip Domain-Driven Whitelist. Menghindari kerentanan eksekusi teks bebas (eval code) dengan abstraksi Token/Enum terenkapsulasi."),
    ("Dual-Channel Accessibility", "Seluruh telegraf ancaman dilengkapi saluran ganda (bentuk piktogram arah + ritme audio), sehingga pemain tidak bergantung hanya pada warna (Colorblind Friendly)."),
    ("CRT Cyber Shader Engine", "Post-processing shader kustom yang mensimulasikan kelengkungan monitor CRT retro, phosphor glow, dan scanlines, memperkuat atmosfer terminal forensik komputer era 90-an."),
    ("Cinematic Outro & Telemetry", "Sistem outro dinamis yang merekap data akurasi, efisiensi eksekusi aturan, dan kondisi integritas, disampaikan melalui dialog typewriter audio-synced bersama A.I.D.A.")
]

for row_idx, (col0, col1) in enumerate(data_table, start=1):
    c0 = table.cell(row_idx, 0)
    c1 = table.cell(row_idx, 1)
    c0.width = col_widths[0]
    c1.width = col_widths[1]
    bg_color = "F1F5F9" if row_idx % 2 == 1 else "FFFFFF"
    set_cell_background(c0, bg_color)
    set_cell_background(c1, bg_color)
    set_cell_margins(c0, top=100, bottom=100, left=150, right=150)
    set_cell_margins(c1, top=100, bottom=100, left=150, right=150)
    
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(col0)
    r0.bold = True
    r0.font.name = "Arial"
    r0.font.size = Pt(9)
    r0.font.color.rgb = C_DARK
    
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(col1)
    r1.font.name = "Arial"
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = C_DARK

# ==================== BAB 4 ====================
add_header("BAB IV. METODOLOGI DAN PROSES REKAYASA PERANGKAT LUNAK", level=1)

add_header("4.1 Pola Arsitektur (Software Architecture)", level=2)
add_body("Pengembangan AlgoRhythm mengedepankan prinsip Clean Architecture dan Separation of Concerns (SoC) pada ekosistem Unity:")
add_bullet(" Berisi domain model murni yang bebas dependensi UnityEngine (Core Combat, Rhythm Clock, Rule Data Structures). Bagian ini dapat diuji secara independen melalui automated unit test dengan kecepatan tinggi.", bold_prefix="• Core Pure-C# Layer: ")
add_bullet(" Menggunakan ScriptableObject sebagai data-driven container untuk komposisi wave, audio track profile, dan parameter malware.", bold_prefix="• Scriptable Data Architecture: ")
add_bullet(" Pemisahan ketat antara state mesin dan visual/audio presentation. Status pertahanan diperbarui melalui C# Events terisolasi tanpa coupling langsung ke UI komponen.", bold_prefix="• Decoupled Event System: ")

add_header("4.2 Penjaminan Kualitas (Quality Assurance & Automated Testing)", level=2)
add_body("Proyek ini dilengkapi rangkaian pengujian otomatis (Automated Unit Testing) menggunakan Unity Test Runner (NUnit) di folder Tests/EditMode. Modul-modul kritis seperti verifikasi toleransi ketukan DSP, eksekusi condition-action rule engine, dan integritas damage loop divalidasi dengan test pass 100% sebelum rilis build, memastikan game terbebas dari bug kritis saat dimainkan juri.")

add_header("4.3 Kepatuhan Hak Cipta dan Aset Bebas Royalti (Zero-Plagiarism)", level=2)
add_body("Sesuai instruksi ketat Bab VII Juknis OTN 2026 mengenai Hak Cipta & Orisinalitas Karya:")
add_bullet(" Kode program inti ditulis mandiri oleh peserta dengan struktur terstruktur dan terdokumentasi rapi.", bold_prefix="• Kode Program: ")
add_bullet(" Elemen UI, layout terminal, dan custom shader dikembangkan secara orisinal.", bold_prefix="• Desain Visual & Visual FX: ")
add_bullet(" Efek suara (SFX) sintetis dan komposisi beat musik menggunakan lisensi Creative Commons Zero (CC0) / Public Domain yang bebas royalti tanpa pelanggaran lisensi komersial pihak ketiga.", bold_prefix="• Audio & Musik: ")

# ==================== BAB 5 ====================
add_header("BAB V. MATRIKS KESESUAIAN TEMA DAN KESIMPULAN", level=1)

add_body("Karya ini menjawab tema utama Olimpiade TIK Informatika Nasional 2026 secara komprehensif:")
add_bullet(" Game menuntut kemampuan abstraksi dan pemecahan masalah cepat. Pemain menganalisis pola ancaman, menghitung konsumsi energi, dan menentukan kapan harus mengintervensi secara manual atau mengandalkan sistem otomasi.", bold_prefix="1. Cerdas Berpikir: ")
add_bullet(" Memvisualisasikan konsep percabangan kode (if-else), event listener, dan arsitektur firewall menjadi mekanik perisai interaktif yang atraktif dan menantang.", bold_prefix="2. Kreatif Berkoding: ")
add_bullet(" Menghadirkan A.I.D.A. sebagai mitra analitik defensif. Game menyampaikan pesan edukatif bahwa kecerdasan buatan merupakan instrumen optimasi yang memerlukan pengawasan, etika pembatasan, dan kendali logika manusia secara bijak.", bold_prefix="3. Bijak Menggunakan AI: ")

add_header("5.2 Kesimpulan", level=2)
add_body("AlgoRhythm: CODE IS SHIELD bukan sekadar game hiburan biasa, melainkan media gamifikasi edukasi informatika yang kokoh secara teknis, inovatif secara konsep permainan, dan sarat akan pesan literasi teknologi generasi baru Indonesia.")

output_path = "/home/rewsaqy/AlgoRhythm2D/docs/Penjelasan_Karya_AlgoRhythm_OTN2026.docx"
doc.save(output_path)
print("Expanded docx saved successfully.")
